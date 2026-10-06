from datetime import datetime
from flask import Blueprint, render_template, request, redirect, url_for, flash
import sqlite3
from pathlib import Path
import shutil
import re

admin_catalog = Blueprint("admin_catalog", __name__)
ROOT = Path(__file__).resolve().parent
DB = Path(__file__).resolve().parent / "data" / "chariow.db"

def db():
    c = sqlite3.connect(DB, timeout=15.0)
    c.row_factory = sqlite3.Row
    return c

def normalize_media(value, media_type):
    raw = str(value or '').strip()
    if not raw:
        return ''
    if raw.lower().startswith(('http://', 'https://', 'data:')):
        return raw
    normalized = raw.replace('\\', '/')
    if media_type == 'image':
        if normalized.startswith('/static/'):
            normalized = normalized[8:]
        elif normalized.startswith('static/'):
            normalized = normalized[7:]
        folder = ROOT / 'static' / 'uploads' / 'products' / 'covers'
        prefix = 'uploads/products/covers'
    else:
        marker = 'protected_data/pdfs/'
        low = normalized.lower()
        if marker in low:
            return Path(normalized[low.find(marker) + len(marker):]).name
        if '/' not in normalized:
            return Path(normalized).name
        folder = ROOT / 'protected_data' / 'pdfs'
        prefix = ''
    source = Path(raw)
    if not source.is_absolute():
        candidate = (ROOT / raw).resolve()
        if candidate.is_file():
            source = candidate
    if not source.is_file():
        return normalized
    folder.mkdir(parents=True, exist_ok=True)
    filename = re.sub(r'[<>:"/\\|?*\\x00-\\x1F]', '_', source.name)
    filename = re.sub(r'\s+', '_', filename).strip(' ._') or 'fichier'
    destination = folder / filename
    if not destination.exists():
        shutil.copy2(source, destination)
    elif destination.stat().st_size != source.stat().st_size:
        stem = destination.stem
        suffix = destination.suffix
        n = 2
        while True:
            candidate = folder / f'{stem}_{n}{suffix}'
            if not candidate.exists():
                destination = candidate
                shutil.copy2(source, destination)
                break
            n += 1
    if media_type == 'image':
        return f'{prefix}/{destination.name}'
    return destination.name



@admin_catalog.route("/admin/produits")
def produits():
    q = request.args.get("q","").strip()
    con=db()
    sql="""SELECT p.*, c.name AS category_name
           FROM products p LEFT JOIN categories c ON c.id=p.category_id
           WHERE 1=1"""
    params=[]
    if q:
        sql += " AND (p.name LIKE ? OR COALESCE(p.description,'') LIKE ?)"
        x=f"%{q}%"; params=[x,x]
    sql += " ORDER BY p.id DESC"
    rows=con.execute(sql,params).fetchall()
    categories=con.execute("SELECT * FROM categories ORDER BY name").fetchall()
    con.close()
    return render_template("admin_products.html",products=rows,categories=categories,q=q)

@admin_catalog.route("/admin/produit/nouveau", methods=["GET","POST"])
def nouveau():
    con=db()
    categories=con.execute("SELECT * FROM categories ORDER BY name").fetchall()
    if request.method=="POST":
        name=request.form.get("name","").strip()
        price=float(request.form.get("price") or 0)
        stock=int(request.form.get("stock") or 0)
        category_id=request.form.get("category_id") or None
        description=request.form.get("description","").strip()
        image=request.form.get("image","").strip()
        active=1 if request.form.get("active") else 0
        if not name:
            con.close()
            return "Le nom du produit est obligatoire",400
        con.execute("""INSERT INTO products
            (name,price,stock,active,description,image,category_id)
            VALUES (?,?,?,?,?,?,?)""",
            (name,price,stock,active,description,image,category_id))
        con.commit(); con.close()
        return redirect(url_for("admin_catalog.produits"))
    con.close()
    return render_template("admin_product_form.html",product=None,categories=categories)

@admin_catalog.route("/admin/produit/<int:product_id>/modifier", methods=["GET","POST"])
def modifier(product_id):
    con=db()
    product=con.execute("SELECT * FROM products WHERE id=?",(product_id,)).fetchone()
    categories=con.execute("SELECT * FROM categories ORDER BY name").fetchall()
    if not product:
        con.close(); return "Produit introuvable",404
    if request.method=="POST":
        name=request.form.get("name","").strip()
        price=float(request.form.get("price") or 0)
        stock=int(request.form.get("stock") or 0)
        category_id=request.form.get("category_id") or None
        description=request.form.get("description","").strip()
        image=request.form.get("image","").strip()
        active=1 if request.form.get("active") else 0
        con.execute("""UPDATE products SET name=?,price=?,stock=?,active=?,
                       description=?,image=?,category_id=? WHERE id=?""",
                    (name,price,stock,active,description,image,category_id,product_id))
        con.commit(); con.close()
        return redirect(url_for("admin_catalog.produits"))
    con.close()
    return render_template("admin_product_form.html",product=product,categories=categories)


# === SASTOUKASTORE EXCEL LIVRES 1 ===
@admin_catalog.route("/admin/produits/export-excel")
def export_excel():
    import io
    from flask import send_file
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return "Le module openpyxl est nécessaire. Installez-le avec : python -m pip install openpyxl", 500

    con = db()
    headers = [
        "id", "name", "slug", "price", "currency", "level", "subject",
        "short_description", "description", "category", "image", "pdf_file",
        "featured", "active", "stock"
    ]
    product_columns = {r["name"] for r in con.execute("PRAGMA table_info(products)").fetchall()}
    select_parts = []
    for col in headers:
        if col == "category":
            select_parts.append("c.name AS category")
        elif col in product_columns:
            select_parts.append(f"p.{col} AS {col}")
        else:
            select_parts.append(f"NULL AS {col}")

    rows = con.execute(
        "SELECT " + ", ".join(select_parts) +
        " FROM products p LEFT JOIN categories c ON c.id=p.category_id ORDER BY p.id ASC"
    ).fetchall()
    con.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Livres"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="101827")
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([row[h] for h in headers])

    widths = [10,42,38,12,12,18,24,42,65,25,45,55,12,12,12]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Instructions")
    info.append(["SASTOUKASTORE — Import / Export des livres"])
    info.append(["Colonne", "Utilisation"])
    for a, b in [
        ("id", "ID existant = mise à jour ; vide = création"),
        ("name", "Titre du livre/support PDF — obligatoire"),
        ("price", "Prix"),
        ("currency", "EUR, USD, FCFA, MAD..."),
        ("level", "Niveau"),
        ("subject", "Matière"),
        ("short_description", "Présentation courte"),
        ("description", "Description détaillée"),
        ("category", "Nom de catégorie ; créée automatiquement si absente"),
        ("image", "URL ou chemin de couverture"),
        ("pdf_file", "Chemin/URL du PDF ; le fichier PDF n'est pas incorporé dans Excel"),
        ("featured", "1 ou 0"),
        ("active", "1 ou 0"),
        ("stock", "Stock ; 999999 recommandé pour un PDF"),
    ]:
        info.append([a, b])
    info.column_dimensions["A"].width = 25
    info.column_dimensions["B"].width = 100

    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    return send_file(
        data, as_attachment=True,
        download_name=f"SASTOUKASTORE_LIVRES_{datetime.now():%Y%m%d_%H%M}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@admin_catalog.route("/admin/produits/modele-excel")
def modele_excel():
    import io
    from flask import send_file
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return "Le module openpyxl est nécessaire. Installez-le avec : python -m pip install openpyxl", 500

    wb = Workbook()
    ws = wb.active
    ws.title = "Livres"
    headers = [
        "id", "name", "slug", "price", "currency", "level", "subject",
        "short_description", "description", "category", "image", "pdf_file",
        "featured", "active", "stock"
    ]
    ws.append(headers)
    ws.append([
        "", "Mathématiques — Additions niveau 1", "", 1.50, "EUR",
        "Primaire", "Mathématiques", "Exercices d'addition pour enfants.",
        "Livre PDF avec exercices progressifs.", "Mathématiques", "",
        "livre_additions_niveau1.pdf", 0, 1, 999999
    ])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="101827")
    for i, width in enumerate([10,42,38,12,12,18,24,42,65,25,45,55,12,12,12], 1):
        ws.column_dimensions[chr(64+i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Instructions")
    info.append(["SASTOUKASTORE", "Modèle d'importation massive des livres"])
    info.append(["Création", "Laisser id vide"])
    info.append(["Mise à jour", "Indiquer l'id existant"])
    info.append(["PDF", "Excel ne transporte pas le PDF ; renseigner pdf_file avec son chemin/référence"])
    info.append(["Catégorie", "Créée automatiquement si elle n'existe pas"])

    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    return send_file(
        data, as_attachment=True,
        download_name="SASTOUKASTORE_MODELE_IMPORT_LIVRES.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@admin_catalog.route("/admin/produits/import-excel", methods=["POST"])
def import_excel():
    from flask import flash
    from werkzeug.utils import secure_filename

    upload = request.files.get("excel_file")
    if not upload or not upload.filename:
        flash("Veuillez sélectionner un fichier Excel (.xlsx).", "error")
        return redirect(url_for("admin_catalog.produits"))
    if not upload.filename.lower().endswith(".xlsx"):
        flash("Format refusé. Utilisez un fichier .xlsx.", "error")
        return redirect(url_for("admin_catalog.produits"))

    try:
        from openpyxl import load_workbook
    except ImportError:
        flash("openpyxl manque. Installez-le avec : python -m pip install openpyxl", "error")
        return redirect(url_for("admin_catalog.produits"))

    temp_dir = ROOT / "data" / "imports"
    temp_dir.mkdir(parents=True, exist_ok=True)
    temp_path = temp_dir / (secure_filename(upload.filename) or "import.xlsx")
    upload.save(temp_path)

    con = db()
    created = updated = skipped = 0
    errors = []

    try:
        wb = load_workbook(temp_path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        raw_headers = list(next(ws.iter_rows(values_only=True)))
        headers = [str(v).strip().lower() if v is not None else "" for v in raw_headers]

        if "name" not in headers:
            raise ValueError("La colonne obligatoire 'name' est absente.")

        def value(row, key, default=""):
            if key not in headers:
                return default
            idx = headers.index(key)
            return row[idx] if idx < len(row) else default

        product_columns = {r["name"] for r in con.execute("PRAGMA table_info(products)").fetchall()}
        writable = {
            c for c in [
                "name","slug","description","price","image","stock","active",
                "currency","level","subject","short_description","featured",
                "product_type","pdf_file"
            ] if c in product_columns
        }

        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            try:
                name = str(value(row, "name", "") or "").strip()
                if not name:
                    skipped += 1
                    continue

                try:
                    price = float(value(row, "price", 0) or 0)
                except (TypeError, ValueError):
                    price = 0.0
                try:
                    stock = int(float(value(row, "stock", 999999) or 999999))
                except (TypeError, ValueError):
                    stock = 999999

                def bool_value(key, default=1):
                    raw = value(row, key, default)
                    if isinstance(raw, bool):
                        return int(raw)
                    txt = str(raw).strip().lower()
                    if txt in {"0","false","non","no","n"}: return 0
                    if txt in {"1","true","oui","yes","y"}: return 1
                    return default

                active = bool_value("active", 1)
                featured = bool_value("featured", 0)

                category_id = None
                category_name = str(value(row, "category", "") or "").strip()
                if category_name and "category_id" in product_columns:
                    cat = con.execute(
                        "SELECT id FROM categories WHERE lower(name)=lower(?) LIMIT 1",
                        (category_name,)
                    ).fetchone()
                    if cat:
                        category_id = cat["id"]
                    else:
                        slug_cat = "".join(ch if ch.isalnum() else "-" for ch in category_name.lower())
                        while "--" in slug_cat:
                            slug_cat = slug_cat.replace("--", "-")
                        slug_cat = slug_cat.strip("-") or "categorie"
                        try:
                            con.execute(
                                "INSERT INTO categories(name,slug,active) VALUES(?,?,1)",
                                (category_name, slug_cat)
                            )
                            category_id = con.execute(
                                "SELECT id FROM categories WHERE lower(name)=lower(?) LIMIT 1",
                                (category_name,)
                            ).fetchone()["id"]
                        except Exception:
                            cat = con.execute(
                                "SELECT id FROM categories WHERE lower(name)=lower(?) LIMIT 1",
                                (category_name,)
                            ).fetchone()
                            category_id = cat["id"] if cat else None

                product = None
                raw_id = value(row, "id", "")
                if raw_id not in ("", None):
                    try:
                        product = con.execute(
                            "SELECT * FROM products WHERE id=?", (int(float(raw_id)),)
                        ).fetchone()
                    except (TypeError, ValueError):
                        pass

                slug = str(value(row, "slug", "") or "").strip()
                if not product and slug and "slug" in product_columns:
                    product = con.execute(
                        "SELECT * FROM products WHERE slug=? LIMIT 1", (slug,)
                    ).fetchone()

                if not slug:
                    slug = "".join(ch if ch.isalnum() else "-" for ch in name.lower())
                    while "--" in slug:
                        slug = slug.replace("--", "-")
                    slug = slug.strip("-") or "livre"

                if not product and "slug" in product_columns:
                    base_slug = slug
                    n = 2
                    while con.execute("SELECT 1 FROM products WHERE slug=?", (slug,)).fetchone():
                        slug = f"{base_slug}-{n}"
                        n += 1

                values = {
                    "name": name,
                    "slug": slug,
                    "price": price,
                    "stock": stock,
                    "active": active,
                    "description": str(value(row, "description", "") or ""),
                    "image": normalize_media(value(row, "image", ""), "image"),
                    "currency": str(value(row, "currency", "EUR") or "EUR"),
                    "level": str(value(row, "level", "") or ""),
                    "subject": str(value(row, "subject", "") or ""),
                    "short_description": str(value(row, "short_description", "") or ""),
                    "featured": featured,
                    "product_type": "digital_pdf",
                    "pdf_file": normalize_media(value(row, "pdf_file", ""), "pdf"),
                }

                if product:
                    assignments = []
                    params = []
                    for col, val in values.items():
                        if col in writable:
                            assignments.append(f"{col}=?")
                            params.append(val)
                    if "category_id" in product_columns:
                        assignments.append("category_id=?")
                        params.append(category_id)
                    params.append(product["id"])
                    con.execute(
                        "UPDATE products SET " + ", ".join(assignments) + " WHERE id=?",
                        params
                    )
                    updated += 1
                else:
                    cols = []
                    vals = []
                    marks = []
                    for col, val in values.items():
                        if col in writable:
                            cols.append(col); vals.append(val); marks.append("?")
                    if "category_id" in product_columns:
                        cols.append("category_id"); vals.append(category_id); marks.append("?")
                    con.execute(
                        "INSERT INTO products (" + ",".join(cols) + ") VALUES (" + ",".join(marks) + ")",
                        vals
                    )
                    created += 1

            except Exception as exc:
                skipped += 1
                errors.append(f"Ligne {row_number}: {exc}")

        con.commit()
        wb.close()
    except Exception as exc:
        con.rollback()
        errors.append(str(exc))
    finally:
        con.close()
        try:
            temp_path.unlink(missing_ok=True)
        except Exception:
            pass

    msg = f"Import Excel terminé : {created} créé(s), {updated} mis à jour, {skipped} ignoré(s)."
    if errors:
        msg += " " + " | ".join(errors[:5])
    flash(msg, "success" if not errors else "warning")
    return redirect(url_for("admin_catalog.produits"))

# === FIN SASTOUKASTORE EXCEL LIVRES 1 ===


# === SASTOUKASTORE DELETE PRODUCTS 2 ===
def _product_has_orders(con, product_id):
    try:
        return int(con.execute(
            "SELECT COUNT(*) FROM order_items WHERE product_id=?",
            (product_id,)
        ).fetchone()[0] or 0)
    except sqlite3.OperationalError:
        return 0


def _delete_product_media(product):
    for field, base in (
        ("image", ROOT / "static"),
        ("pdf_file", ROOT / "protected_data" / "pdfs"),
    ):
        value = str(product[field] or "").strip()
        if not value:
            continue

        value = value.replace("/", "\\")
        path = Path(value)

        if not path.is_absolute():
            if field == "image":
                path = base / value.lstrip("\\/")
            else:
                path = base / value

        try:
            path = path.resolve()
        except Exception:
            continue

        if path.is_file():
            try:
                path.unlink()
            except Exception:
                pass


def _delete_one_product(con, product_id):
    product = con.execute(
        "SELECT id, name, image, pdf_file FROM products WHERE id=?",
        (product_id,)
    ).fetchone()

    if not product:
        return "missing", ""

    if _product_has_orders(con, product_id) > 0:
        return "protected", str(product["name"])

    try:
        con.execute(
            "DELETE FROM product_variants WHERE product_id=?",
            (product_id,)
        )
    except sqlite3.OperationalError:
        pass

    con.execute(
        "DELETE FROM products WHERE id=?",
        (product_id,)
    )

    _delete_product_media(product)
    return "deleted", str(product["name"])


@admin_catalog.route("/admin/produit/<int:product_id>/supprimer", methods=["POST"])
def supprimer(product_id):
    con = db()
    try:
        status, name = _delete_one_product(con, product_id)

        if status == "deleted":
            con.commit()
            flash(f"Livre supprimé : {name}", "success")
        elif status == "protected":
            con.rollback()
            flash(
                f"Suppression refusée : « {name} » est déjà utilisé dans une commande.",
                "warning"
            )
        else:
            con.rollback()
            flash("Livre introuvable.", "warning")

    except Exception as exc:
        con.rollback()
        flash(f"Suppression impossible : {exc}", "error")
    finally:
        con.close()

    return redirect(url_for("admin_catalog.produits"))


@admin_catalog.route("/admin/produits/supprimer-masse", methods=["POST"])
def supprimer_masse():
    raw_ids = request.form.getlist("product_ids")
    ids = []

    for raw in raw_ids:
        try:
            pid = int(raw)
            if pid > 0:
                ids.append(pid)
        except (TypeError, ValueError):
            pass

    ids = list(dict.fromkeys(ids))

    if not ids:
        flash("Aucun livre sélectionné.", "warning")
        return redirect(url_for("admin_catalog.produits"))

    con = db()
    deleted = []
    protected = []
    missing = []

    try:
        for pid in ids:
            status, name = _delete_one_product(con, pid)

            if status == "deleted":
                deleted.append(name)
            elif status == "protected":
                protected.append(name)
            else:
                missing.append(pid)

        con.commit()

        summary = []
        if deleted:
            summary.append(f"{len(deleted)} supprimé(s)")
        if protected:
            summary.append(f"{len(protected)} protégé(s)")
        if missing:
            summary.append(f"{len(missing)} introuvable(s)")

        flash(
            "Suppression multiple : " + ", ".join(summary) + ".",
            "success" if deleted and not protected else "warning"
        )

        if protected:
            flash(
                "Protégés car déjà vendus : " + " | ".join(protected[:8]),
                "warning"
            )

    except Exception as exc:
        con.rollback()
        flash(f"Suppression multiple impossible : {exc}", "error")
    finally:
        con.close()

    return redirect(url_for("admin_catalog.produits"))

# === FIN SASTOUKASTORE DELETE PRODUCTS 2 ===

@admin_catalog.route("/admin/produit/<int:product_id>/stock", methods=["POST"])
def stock(product_id):
    value=int(request.form.get("stock") or 0)
    if value < 0: value=0
    con=db()
    con.execute("UPDATE products SET stock=? WHERE id=?",(value,product_id))
    con.commit(); con.close()
    return redirect(url_for("admin_catalog.produits"))

@admin_catalog.route("/admin/produit/<int:product_id>/toggle", methods=["POST"])
def toggle(product_id):
    con=db()
    con.execute("""UPDATE products SET active=CASE WHEN active=1 THEN 0 ELSE 1 END
                   WHERE id=?""",(product_id,))
    con.commit(); con.close()
    return redirect(url_for("admin_catalog.produits"))

# === SASTOUKASTORE ADMIN DELETE BOOK V3 ===
def _ssk_v3_book_has_purchases(con, product_id):
    try:
        row = con.execute(
            "SELECT COUNT(*) FROM order_items WHERE product_id=?",
            (product_id,)
        ).fetchone()
        return int(row[0] or 0)
    except sqlite3.OperationalError:
        return 0


def _ssk_v3_remove_media(product):
    for field, base in (
        ("image", ROOT / "static"),
        ("pdf_file", ROOT / "protected_data" / "pdfs"),
    ):
        try:
            raw = str(product[field] or "").strip()
        except Exception:
            raw = ""

        if not raw:
            continue

        if raw.lower().startswith(("http://", "https://")):
            continue

        candidate = Path(raw.replace("/", "\\"))
        if not candidate.is_absolute():
            candidate = base / raw.lstrip("\\/")

        try:
            candidate = candidate.resolve()
        except Exception:
            continue

        try:
            if field == "image":
                candidate.relative_to((ROOT / "static").resolve())
            else:
                candidate.relative_to(
                    (ROOT / "protected_data" / "pdfs").resolve()
                )
        except Exception:
            continue

        if candidate.is_file():
            try:
                candidate.unlink()
            except Exception:
                pass


@admin_catalog.route(
    "/admin/livre/<int:product_id>/supprimer-definitivement",
    methods=["POST"]
)
def supprimer_livre_definitivement_v3(product_id):
    con = db()
    try:
        product = con.execute(
            "SELECT id, name, image, pdf_file FROM products WHERE id=?",
            (product_id,)
        ).fetchone()

        if not product:
            con.rollback()
            flash("Livre introuvable.", "warning")
            return redirect(url_for("admin_catalog.produits"))

        if _ssk_v3_book_has_purchases(con, product_id) > 0:
            con.rollback()
            flash(
                f"Suppression refusée : « {product['name']} » est déjà "
                "présent dans les achats. Le livre est conservé.",
                "warning"
            )
            return redirect(url_for("admin_catalog.produits"))

        try:
            con.execute(
                "DELETE FROM product_variants WHERE product_id=?",
                (product_id,)
            )
        except sqlite3.OperationalError:
            pass

        con.execute(
            "DELETE FROM products WHERE id=?",
            (product_id,)
        )
        con.commit()

        _ssk_v3_remove_media(product)
        flash(f"Livre supprimé : {product['name']}", "success")

    except Exception as exc:
        con.rollback()
        flash(f"Suppression impossible : {exc}", "error")
    finally:
        con.close()

    return redirect(url_for("admin_catalog.produits"))

# === FIN SASTOUKASTORE ADMIN DELETE BOOK V3 ===
